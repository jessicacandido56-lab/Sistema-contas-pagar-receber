from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import date
def criar_planilha_contas_pagar():
    workbook = Workbook()
    planilha = workbook.active

    planilha.title = "Contas a Pagar"
    azul = "4F81BD"
    branco = "FFFFFF"
    planilha.merge_cells("A1:E1")
    planilha["A1"] = "AGOSTO 2026"

    planilha.merge_cells("A2:E2")
    planilha["A2"] = "CONTAS À PAGAR"
    for linha in [1, 2]:
        celula = planilha[f"A{linha}"]
        celula.fill = PatternFill(
            fill_type="solid",
            fgColor=azul
        )
        celula.font = Font(
            color=branco,
            bold=True,
            size=14
        )
        celula.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )
    cabecalhos = ["DIA", "DATA", "FORNECEDOR", "DOCUMENTO", "VALOR"]

    for coluna, titulo in enumerate(cabecalhos, start=1):
        planilha.cell(row=3, column=coluna, value=titulo)
    

    for celula in planilha[3]:
        celula.fill = PatternFill(
            fill_type="solid",
            fgColor=azul
        )
        celula.font = Font(
            color=branco,
            bold=True
        )
        celula.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )
    planilha.column_dimensions["A"].width = 15
    planilha.column_dimensions["B"].width = 15
    planilha.column_dimensions["C"].width = 40
    planilha.column_dimensions["D"].width = 20
    planilha.column_dimensions["E"].width = 18
    dias_semana = [
        "SEGUNDA", "TERÇA", "QUARTA",
        "QUINTA", "SEXTA", "SÁBADO", "DOMINGO"
    ]

    linha = 4

    for dia in range(1, 32):
        data_atual = date(2026, 8, dia)
        nome_dia = dias_semana[data_atual.weekday()]

        for _ in range(4):
            planilha.cell(
                row=linha,
                column=1,
                value=nome_dia
            )
            planilha.cell(
                row=linha,
                column=5
            ).number_format = 'R$ #,##0.00'

            planilha.cell(
                row=linha,
                column=2,
                value=data_atual
            )

            planilha.cell(
                row=linha,
                column=2
            ).number_format = "dd/mm/yyyy"

            linha += 1

        planilha.cell(
            row=linha,
            column=1,
            value=nome_dia
        )

        planilha.cell(
            row=linha,
            column=2,
            value=data_atual
        )

        planilha.cell(
            row=linha,
            column=3,
            value="- TOTAL"
        )
        primeira_linha = linha - 4
        ultima_linha = linha - 1

        planilha.cell(
            row=linha,
            column=5,
            value=f"=SUM(E{primeira_linha}:E{ultima_linha})"
        )

        planilha.cell(
            row=linha,
            column=5
        ).number_format = 'R$ #,##0.00'
        amarelo = "FFF2CC"

        for coluna in range(1, 6):
            celula_total = planilha.cell(
                row=linha,
                column=coluna
            )

            celula_total.fill = PatternFill(
                fill_type="solid",
                fgColor=amarelo
            )

            celula_total.font = Font(
                bold=True
            )

        planilha.cell(
            row=linha,
            column=2
        ).number_format = "dd/mm/yyyy"

        linha += 1
    workbook.save("relatorios/contas_a_pagar.xlsx")
criar_planilha_contas_pagar()